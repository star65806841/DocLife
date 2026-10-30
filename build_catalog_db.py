#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_catalog_db.py
"""

import sys
import os
import json
import sqlite3
import zipfile
import argparse
import traceback
import re

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

print("=== build_catalog_db.py START ===", flush=True)

try:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    print("[OK] openpyxl imported", flush=True)
except ImportError as e:
    print(f"[FATAL] openpyxl not installed: {e}", flush=True)
    sys.exit(1)

# ============================================================================
# Config & Schema
# ============================================================================
DISPLAY_NAME_MAP = {
    "examples":    "示例表格合集",
    "sichuan":     "四川省",
    "guangdong":   "广东省",
    "tunnel":      "隧道工程质检用表",
    "sub_project": "分部工程",
    "item_project": "分项工程",
    "tunnel_qualityinspectionform": "隧道工程质量检验评定表",
}

VIRTUAL_ROW_LIMIT = 5000
VIRTUAL_COL_LIMIT = 500

def display_name_for(name):
    return DISPLAY_NAME_MAP.get(name, name)

SCHEMA_SQL = """
CREATE TABLE catalogs (id INTEGER PRIMARY KEY, code TEXT UNIQUE NOT NULL, display_name TEXT NOT NULL);
CREATE TABLE folders (id INTEGER PRIMARY KEY, catalog_id INTEGER NOT NULL, key TEXT NOT NULL, display_name TEXT NOT NULL, sort_order INTEGER DEFAULT 0);
CREATE TABLE templates (id INTEGER PRIMARY KEY, folder_id INTEGER NOT NULL, code TEXT NOT NULL, title TEXT NOT NULL, json_path TEXT, pdf_path TEXT, xlsx_path TEXT, sort_order INTEGER DEFAULT 0);
CREATE TABLE template_sheets (
    id INTEGER PRIMARY KEY, template_id INTEGER NOT NULL, page INTEGER NOT NULL, name TEXT NOT NULL, title TEXT NOT NULL, sort_order INTEGER DEFAULT 0,
    row_count INTEGER DEFAULT 0, col_count INTEGER DEFAULT 0, default_row_height REAL DEFAULT 14.25, default_col_width REAL DEFAULT 8.43,
    page_orientation INTEGER DEFAULT 0, page_size TEXT DEFAULT 'A4', margin_top REAL DEFAULT 15.0, margin_bottom REAL DEFAULT 15.0, margin_left REAL DEFAULT 15.0, margin_right REAL DEFAULT 15.0
);
CREATE TABLE sheet_cells (
    sheet_id INTEGER NOT NULL, row INTEGER NOT NULL, col INTEGER NOT NULL, row_span INTEGER DEFAULT 1, col_span INTEGER DEFAULT 1,
    value TEXT DEFAULT '', formula TEXT, cell_type INTEGER DEFAULT 0, data_key TEXT,
    font_family TEXT DEFAULT '宋体', font_size REAL DEFAULT 10.5, font_bold INTEGER DEFAULT 0, font_italic INTEGER DEFAULT 0, font_underline INTEGER DEFAULT 0,
    font_color TEXT DEFAULT '#000000', bg_color TEXT DEFAULT '#FFFFFF', h_align INTEGER DEFAULT 0, v_align INTEGER DEFAULT 1,
    border_top INTEGER DEFAULT 0, border_bottom INTEGER DEFAULT 0, border_left INTEGER DEFAULT 0, border_right INTEGER DEFAULT 0, wrap_text INTEGER DEFAULT 1,
    PRIMARY KEY (sheet_id, row, col)
);
CREATE TABLE sheet_row_heights (sheet_id INTEGER NOT NULL, row INTEGER NOT NULL, height REAL NOT NULL, PRIMARY KEY (sheet_id, row));
CREATE TABLE sheet_col_widths (sheet_id INTEGER NOT NULL, col INTEGER NOT NULL, width REAL NOT NULL, PRIMARY KEY (sheet_id, col));
"""

# ============================================================================
# Style helpers
# ============================================================================
def color_to_hex(color):
    if color is None: return None
    t = getattr(color, "type", None)
    if t == "rgb" and isinstance(color.rgb, str):
        s = color.rgb
        if len(s) == 8: return "#" + s[2:]
        if len(s) == 6: return "#" + s
    return None

def border_style_to_int(style):
    return {"thin": 1, "medium": 2, "thick": 3}.get(style, 0)

def parse_cell_style(cell):
    style = {}
    has_non_default = False
    try:
        f = cell.font
        if f:
            if f.name: 
                style["font_family"] = f.name
            if f.sz: 
                style["font_size"] = float(f.sz)
                has_non_default = True
            style["font_bold"]      = 1 if f.bold else 0
            style["font_italic"]    = 1 if f.italic else 0
            style["font_underline"] = 1 if f.underline else 0
            fg = color_to_hex(f.color)
            if fg: 
                style["font_color"] = fg
            if f.bold or f.italic or f.underline or fg: 
                has_non_default = True
    except: pass

    try:
        fill = cell.fill
        if fill and fill.fill_type == "solid":
            bg = color_to_hex(fill.start_color)
            if bg and bg not in ("#FFFFFF", "#000000"):
                style["bg_color"] = bg
                has_non_default = True
    except: pass

    try:
        a = cell.alignment
        if a:
            if a.horizontal == "center": style["h_align"] = 1; has_non_default = True
            elif a.horizontal == "right": style["h_align"] = 2; has_non_default = True
            if a.vertical == "top": style["v_align"] = 0; has_non_default = True
            elif a.vertical == "bottom": style["v_align"] = 2; has_non_default = True
            if a.wrap_text is False: style["wrap_text"] = 0
    except: pass

    try:
        b = cell.border
        if b:
            top    = border_style_to_int(b.top.style    if b.top    else None)
            bottom = border_style_to_int(b.bottom.style if b.bottom else None)
            left   = border_style_to_int(b.left.style   if b.left   else None)
            right  = border_style_to_int(b.right.style  if b.right  else None)
            if top or bottom or left or right:
                style["border_top"] = top
                style["border_bottom"] = bottom
                style["border_left"] = left
                style["border_right"] = right
                has_non_default = True
    except: pass
    return style if has_non_default else None

# ============================================================================
# 列宽解析
# ============================================================================
def get_real_col_widths(xlsx_path, sheet_idx):
    widths = {}
    custom_widths = {}
    try:
        with zipfile.ZipFile(xlsx_path, 'r') as zf:
            xml_path = f"xl/worksheets/sheet{sheet_idx + 1}.xml"
            if xml_path not in zf.namelist():
                possible = [n for n in zf.namelist() if n.startswith("xl/worksheets/sheet")]
                if sheet_idx < len(possible): xml_path = possible[sheet_idx]
                else: return widths

            with zf.open(xml_path) as f:
                content = f.read().decode('utf-8', errors='ignore')
                cols_matches = re.finditer(r'<(?:\w+:)?col\b[^>]*>', content, re.IGNORECASE)
                for col_match in cols_matches:
                    tag = col_match.group(0)
                    min_m = re.search(r'min=["\'](\d+)["\']', tag)
                    max_m = re.search(r'max=["\'](\d+)["\']', tag)
                    w_m = re.search(r'width=["\']([\d.]+)["\']', tag)
                    cw_m = re.search(r'customWidth=["\']1["\']', tag)
                    
                    if min_m and max_m and w_m:
                        min_c = int(min_m.group(1))
                        max_c = int(max_m.group(1))
                        width = float(w_m.group(1))
                        is_custom = cw_m is not None
                        
                        for c in range(min_c, max_c + 1):
                            if is_custom:
                                custom_widths[c] = width
                            elif c not in custom_widths:
                                widths[c] = width
                                
        for c, w in custom_widths.items():
            widths[c] = w
    except Exception as e:
        print(f"    [WARN] XML parse error: {e}", flush=True)
    return widths

# ============================================================================
# Effective bounds
# ============================================================================
def effective_bounds(ws):
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_col > VIRTUAL_COL_LIMIT: max_col = VIRTUAL_COL_LIMIT
    if max_row > VIRTUAL_ROW_LIMIT:
        real_max = 0
        for r in range(max_row, 0, -1):
            for row_cells in ws.iter_rows(min_row=r, max_row=r, min_col=1, max_col=max_col):
                if any(not isinstance(c, MergedCell) and c.value is not None for c in row_cells):
                    real_max = r
                    break
            if real_max > 0: break
        max_row = real_max
    return max_row, max_col

# ============================================================================
# xlsx -> db
# ============================================================================
def convert_xlsx_to_db(cur, xlsx_path, sheet_rows):
    print(f"  -> parsing xlsx: {os.path.basename(xlsx_path)}", flush=True)
    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        print(f"[WARN] cannot load xlsx {xlsx_path}: {e}", flush=True)
        return

    ws_list = wb.worksheets
    sheet_map = {}
    used = set()
    for i, ws in enumerate(ws_list):
        target = None
        for sid, name, _title in sheet_rows:
            if sid not in used and name == ws.title:
                target = sid
                break
        if target is None:
            for sid, name, _title in sheet_rows:
                if sid not in used:
                    target = sid
                    break
        if target is not None:
            sheet_map[i] = target
            used.add(target)

    for ws_idx, sheet_id in sheet_map.items():
        ws = ws_list[ws_idx]
        print(f"     [{ws_idx+1}/{len(ws_list)}] '{ws.title}' …", flush=True)

        max_row, max_col = effective_bounds(ws)
        if max_row == 0 or max_col == 0: continue

        real_def_col_width = ws.sheet_format.defaultColWidth or 8.43
        real_def_row_height = ws.sheet_format.defaultRowHeight or 14.25

        cur.execute(
            "UPDATE template_sheets SET row_count=?, col_count=?, default_row_height=?, default_col_width=? WHERE id=?",
            (max_row, max_col, real_def_row_height, real_def_col_width, sheet_id))

        # 行高
        merged_rows = set()
        for m in ws.merged_cells.ranges:
            for r_idx in range(m.min_row, m.max_row + 1): merged_rows.add(r_idx)

        for r in range(1, max_row + 1):
            dim = ws.row_dimensions.get(r)
            height = real_def_row_height
            if dim:
                if dim.customHeight and dim.height: height = dim.height
                elif dim.height:
                    if r in merged_rows: height = real_def_row_height
                    else: height = dim.height
            cur.execute("INSERT OR REPLACE INTO sheet_row_heights VALUES (?,?,?)", (sheet_id, r - 1, float(height)))

        # 列宽
        real_widths = get_real_col_widths(xlsx_path, ws_idx)
        for c in range(1, max_col + 1):
            width = real_widths.get(c, real_def_col_width)
            cur.execute("INSERT OR REPLACE INTO sheet_col_widths VALUES (?,?,?)", (sheet_id, c - 1, float(width)))

        # 合并单元格
        merges = {}
        for m in ws.merged_cells.ranges:
            merges[(m.min_row - 1, m.min_col - 1)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)

        # 单元格内容
        cell_count = 0
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                if isinstance(cell, MergedCell): continue
                r0, c0 = cell.row - 1, cell.column - 1
                
                value, formula = "", None
                if cell.value is not None:
                    if isinstance(cell.value, str) and cell.value.startswith("="): formula = cell.value
                    else: value = str(cell.value)

                style = parse_cell_style(cell)
                is_merge_origin = (r0, c0) in merges
                if not value and not formula and not style and not is_merge_origin: continue

                row_span, col_span = merges.get((r0, c0), (1, 1))
                style = style or {}

                # ★ 强力诊断日志：打印第 4、5 行（r0=3,4）有内容的单元格字体
                if r0 in (3, 4) and value.strip():
                    print(f"       [DEBUG FONT] 行{r0} 列{c0} 值:'{value[:5]}' -> openpyxl读取的字号: {cell.font.sz}")

                cur.execute(
                    "INSERT OR REPLACE INTO sheet_cells "
                    "(sheet_id,row,col,row_span,col_span,value,formula,font_family,font_size,font_bold,font_italic,font_underline,font_color,bg_color,h_align,v_align,border_top,border_bottom,border_left,border_right,wrap_text) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (sheet_id, r0, c0, row_span, col_span, value, formula,
                     style.get("font_family", "宋体"), style.get("font_size", 10.5), style.get("font_bold", 0),
                     style.get("font_italic", 0), style.get("font_underline", 0), style.get("font_color", "#000000"),
                     style.get("bg_color", "#FFFFFF"), style.get("h_align", 0), style.get("v_align", 1),
                     style.get("border_top", 0), style.get("border_bottom", 0), style.get("border_left", 0),
                     style.get("border_right", 0), style.get("wrap_text", 1)))
                cell_count += 1
        print(f"       done, {cell_count} cells", flush=True)

# ============================================================================
# Build db & zip
# ============================================================================
def build_db(collection_dir, collection_code, display_name, output_db):
    if not os.path.isdir(collection_dir): return 1
    if os.path.exists(output_db): os.remove(output_db)
    conn = sqlite3.connect(output_db)
    cur = conn.cursor()
    cur.executescript(SCHEMA_SQL)
    cur.execute("INSERT INTO catalogs (code, display_name) VALUES (?, ?)", (collection_code, display_name))
    catalog_id = cur.lastrowid

    folder_order = 0
    n_tpl = n_sheet = n_cell = 0

    for folder_name in sorted(os.listdir(collection_dir)):
        folder_path = os.path.join(collection_dir, folder_name)
        if not os.path.isdir(folder_path): continue
        cur.execute("INSERT INTO folders (catalog_id, key, display_name, sort_order) VALUES (?,?,?,?)", (catalog_id, folder_name, display_name_for(folder_name), folder_order))
        folder_id = cur.lastrowid
        folder_order += 1

        tpl_order = 0
        for json_file in sorted(os.listdir(folder_path)):
            if not json_file.endswith(".json"): continue
            json_full = os.path.join(folder_path, json_file)
            try:
                with open(json_full, "r", encoding="utf-8") as f: data = json.load(f)
            except Exception as e: continue

            base = os.path.splitext(json_file)[0]
            pdf_rel, xlsx_rel = f"{folder_name}/{base}.pdf", f"{folder_name}/{base}.xlsx"

            cur.execute("INSERT INTO templates (folder_id, code, title, json_path, pdf_path, xlsx_path, sort_order) VALUES (?,?,?,?,?,?,?)",
                (folder_id, base, data.get("title", base), f"{folder_name}/{json_file}",
                 pdf_rel if os.path.exists(os.path.join(collection_dir, pdf_rel)) else None,
                 xlsx_rel if os.path.exists(os.path.join(collection_dir, xlsx_rel)) else None, tpl_order))
            template_id = cur.lastrowid
            tpl_order += 1
            n_tpl += 1

            sheet_rows = []
            for i, sheet in enumerate(data.get("sheets", [])):
                cur.execute("INSERT INTO template_sheets (template_id, page, name, title, sort_order) VALUES (?,?,?,?,?)",
                    (template_id, int(sheet.get("page", i)), str(sheet.get("name", "")), str(sheet.get("title", "")), i))
                sheet_rows.append((cur.lastrowid, str(sheet.get("name", "")), str(sheet.get("title", ""))))
                n_sheet += 1

            xlsx_full = os.path.join(collection_dir, xlsx_rel)
            if os.path.exists(xlsx_full):
                before = cur.execute("SELECT COUNT(*) FROM sheet_cells").fetchone()[0]
                convert_xlsx_to_db(cur, xlsx_full, sheet_rows)
                after = cur.execute("SELECT COUNT(*) FROM sheet_cells").fetchone()[0]
                n_cell += (after - before)

    conn.commit()
    conn.close()
    print(f"[OK] db: {output_db}\n     templates={n_tpl}  sheets={n_sheet}  cells={n_cell}", flush=True)
    return 0

def build_zip(collection_dir, output_zip, extra_files):
    if os.path.exists(output_zip): os.remove(output_zip)
    skip = {".DS_Store", "Thumbs.db", "desktop.ini"}
    files = []
    for root, dirs, names in os.walk(collection_dir):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for name in names:
            if name in skip or name.startswith("."): continue
            full = os.path.join(root, name)
            files.append((full, os.path.relpath(full, collection_dir).replace(os.sep, "/")))
    files.extend(extra_files)
    with zipfile.ZipFile(output_zip, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for full, rel in files:
            info = zipfile.ZipInfo.from_file(full, rel)
            info.flag_bits |= 0x800
            info.external_attr = (0o644 << 16)
            with open(full, "rb") as f: zf.writestr(info, f.read())

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--collection", required=True)
    ap.add_argument("--src-root", default="resources/data")
    args = ap.parse_args()

    coll_dir = os.path.join(os.path.abspath(args.src_root), args.collection)
    if build_db(coll_dir, args.collection, display_name_for(args.collection), os.path.join(coll_dir, "catalog.db")) != 0:
        return 1
    build_zip(coll_dir, os.path.join(os.path.abspath(args.src_root), f"{args.collection}.zip"), [(os.path.join(coll_dir, "catalog.db"), "catalog.db")])
    print("=== DONE ===", flush=True)

if __name__ == "__main__":
    sys.exit(main())